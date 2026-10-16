from pathlib import Path

from docx import Document

from src.docx_to_wechat import convert, parse_docx, _looks_like_legacy_four_part
from src.weekly_report_to_wechat import convert as convert_weekly_report


def test_convert_formatted_docx_preserves_document_flow(tmp_path: Path):
    docx_path = tmp_path / "article.docx"
    doc = Document()
    doc.add_paragraph("2026年7月信用卡资讯整合点评")
    doc.add_paragraph('含"引号"与小于号<测试>')
    doc.add_heading("本期导读", level=1)
    doc.add_paragraph("这是一段导读正文。")
    doc.add_heading("一、中国银行 · 测试卡", level=2)
    doc.add_heading("卡片定位与核心亮点", level=3)
    doc.add_paragraph("权益核心：AI权益", style="List Paragraph")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "权益"
    table.cell(0, 1).text = "规则"
    table.cell(1, 0).text = "返现"
    table.cell(1, 1).text = "消费满100元且<200元"
    doc.add_heading("专业点评", level=3)
    doc.add_paragraph('点评包含"智能服务"。')
    doc.save(docx_path)

    out, title = convert(docx_path)
    html = out.read_text(encoding="utf-8")

    assert title == "2026年7月信用卡资讯整合点评"
    assert "共 276 条" not in html
    assert html.count("<table") == 1
    assert "<blockquote" in html
    assert "&amp;quot;" not in html
    assert "&amp;lt;" not in html
    assert "&lt;测试&gt;" in html
    assert "消费满100元且&lt;200元" in html


def test_legacy_four_part_detection_unescapes_xml_entities(tmp_path: Path):
    docx_path = tmp_path / "legacy.docx"
    doc = Document()
    for title in ("银行A - 测试卡", "银行B - 测试活动"):
        doc.add_paragraph(title)
        doc.add_paragraph("来源：公众号")
        doc.add_paragraph("核心内容：消费满100元且<200元")
        doc.add_paragraph('点评：关注"门槛"。')
    doc.save(docx_path)

    paras = parse_docx(docx_path)

    assert _looks_like_legacy_four_part(paras)
    assert any("<200元" in p["text"] for p in paras)
    assert any('关注"门槛"' in p["text"] for p in paras)


def test_weekly_converter_falls_back_for_single_article_docx(tmp_path: Path):
    """A non-weekly document must not silently become an HTML shell."""
    docx_path = tmp_path / "single_article.docx"
    doc = Document()
    doc.add_heading("测试信用卡活动", level=1)
    doc.add_heading("核心内容", level=2)
    doc.add_paragraph("测试正文：完成指定消费可享50元立减。")
    doc.save(docx_path)

    out = convert_weekly_report(docx_path)
    rendered = out.read_text(encoding="utf-8")

    assert "测试正文：完成指定消费可享50元立减。" in rendered
    assert rendered.count("<p ") >= 1

